# -*- coding: utf-8 -*-
from odoo import models, fields, api, _
from odoo.exceptions import UserError
import base64
from datetime import datetime
import logging

_logger = logging.getLogger(__name__)


class PharmacyExpenseImportWizard(models.TransientModel):
    _name = 'pharmacy.expense.import.wizard'
    _description = 'Wizard para importar gastos de farmacia desde Excel'

    state = fields.Selection([
        ('step1', 'Carga'),
        ('step2', 'Preview'),
        ('step3', 'Resultado'),
    ], default='step1')

    # Step 1 fields
    farmacia_id = fields.Many2one(
        'sindicato.proveedor',
        string='Farmacia',
        domain=[('tipo', '=', 'farmacia'), ('activo', '=', True)],
    )
    date_month = fields.Selection(
        selection=[
            ('01', 'Enero'), ('02', 'Febrero'), ('03', 'Marzo'),
            ('04', 'Abril'), ('05', 'Mayo'), ('06', 'Junio'),
            ('07', 'Julio'), ('08', 'Agosto'), ('09', 'Septiembre'),
            ('10', 'Octubre'), ('11', 'Noviembre'), ('12', 'Diciembre'),
        ],
        string='Mes',
        required=True,
        default=lambda self: datetime.now().strftime('%m'),
    )

    @api.model
    def _get_year_selection(self):
        current_year = datetime.now().year
        return [(str(y), str(y)) for y in reversed(range(current_year - 5, current_year + 2))]

    date_year = fields.Selection(
        selection=_get_year_selection,
        string='Año',
        required=True,
        default=lambda self: str(datetime.now().year),
    )
    file_data = fields.Binary(string='Archivo Excel', attachment=False)
    file_name = fields.Char(string='Nombre del archivo')

    # Step 2 fields - preview
    preview_line_ids = fields.One2many(
        'pharmacy.expense.import.preview',
        'wizard_id',
        string='Preview',
    )
    total_afiliados = fields.Integer(string='Total afiliados', readonly=True)
    total_tickets = fields.Integer(string='Total tickets', readonly=True)
    total_receta = fields.Float(string='Total Bajo Receta', readonly=True)
    total_venta_libre = fields.Float(string='Total Venta Libre', readonly=True)
    total_general = fields.Float(string='Total General', readonly=True)
    has_overwrites = fields.Boolean(string='Hay sobreescrituras', readonly=True)

    # Step 3 fields - result
    result_message = fields.Text(string='Resultado', readonly=True)

    def action_parse_file(self):
        """Step 1 -> Step 2: Parse the Excel file and show preview."""
        self.ensure_one()
        if not self.file_data:
            raise UserError(_('Debe seleccionar un archivo Excel.'))
        if not self.farmacia_id:
            raise UserError(_('Debe seleccionar una farmacia.'))

        try:
            import openpyxl
            from io import BytesIO
        except ImportError:
            raise UserError(_('La librería openpyxl no está instalada.'))

        file_content = base64.b64decode(self.file_data)
        wb = openpyxl.load_workbook(BytesIO(file_content), data_only=True)
        ws = wb.active

        # Parse rows starting from row 4 (skip header rows 1-3)
        tickets_by_uid = {}  # uid -> list of ticket dicts
        for row in ws.iter_rows(min_row=4, max_row=ws.max_row, values_only=False):
            # Skip empty rows or summary rows (no uid)
            uid_val = row[3].value  # Column D = N° Afiliado
            if uid_val is None:
                continue
            try:
                uid = int(uid_val)
            except (ValueError, TypeError):
                continue

            nombre = row[2].value or ''  # Column C
            fecha_raw = row[5].value     # Column F
            monto_receta = self._parse_number(row[6].value)   # Column G
            monto_venta_libre = self._parse_number(row[7].value)  # Column H

            # Parse date
            fecha = None
            if isinstance(fecha_raw, datetime):
                fecha = fecha_raw.date()
            elif isinstance(fecha_raw, str):
                for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
                    try:
                        fecha = datetime.strptime(fecha_raw, fmt).date()
                        break
                    except ValueError:
                        continue

            # Parse order number
            numero_orden = 0
            try:
                numero_orden = int(row[0].value) if row[0].value else 0
            except (ValueError, TypeError):
                pass

            if uid not in tickets_by_uid:
                tickets_by_uid[uid] = []
            tickets_by_uid[uid].append({
                'uid': uid,
                'nombre': nombre,
                'numero_orden': numero_orden,
                'fecha': fecha,
                'monto_receta': monto_receta,
                'monto_venta_libre': monto_venta_libre,
            })

        if not tickets_by_uid:
            raise UserError(_('No se encontraron datos en el archivo Excel.'))

        # Look up affiliates and check for existing records
        Affiliate = self.env['affiliation.affiliate']
        Expense = self.env['affiliate.pharmacy.expenses']

        # Clear old preview lines
        self.preview_line_ids.unlink()

        preview_lines = []
        total_receta = 0.0
        total_venta_libre = 0.0
        has_overwrites = False
        no_encontrados = []

        # First pass: check all affiliates exist
        for uid, tickets in tickets_by_uid.items():
            affiliate = Affiliate.search([('uid', '=', uid)], limit=1)
            if not affiliate:
                nombre_ref = tickets[0]['nombre'] if tickets else str(uid)
                no_encontrados.append(f"UID {uid} - {nombre_ref}")

        if no_encontrados:
            raise UserError(_(
                'Los siguientes afiliados no fueron encontrados en el sistema. '
                'Deben existir antes de importar:\n\n%s'
            ) % '\n'.join(no_encontrados))

        # Second pass: build preview
        for uid, tickets in tickets_by_uid.items():
            affiliate = Affiliate.search([('uid', '=', uid)], limit=1)

            sum_receta = sum(t['monto_receta'] for t in tickets)
            sum_venta_libre = sum(t['monto_venta_libre'] for t in tickets)

            # Check if expense record exists
            expense = Expense.search([
                ('affiliate_id', '=', affiliate.id),
                ('month', '=', self.date_month),
                ('year', '=', int(self.date_year)),
            ], limit=1)

            overwrite = False
            if expense:
                existing_line = expense.linea_gastos_ids.filtered(
                    lambda l: l.farmacia_id.id == self.farmacia_id.id
                )
                if existing_line and (existing_line.gasto_plan != 0 or existing_line.gasto_venta_libre != 0):
                    overwrite = True
                    has_overwrites = True

            preview_lines.append((0, 0, {
                'affiliate_id': affiliate.id,
                'uid': uid,
                'nombre_archivo': tickets[0]['nombre'],
                'cantidad_tickets': len(tickets),
                'total_receta': sum_receta,
                'total_venta_libre': sum_venta_libre,
                'sobreescribe': overwrite,
            }))
            total_receta += sum_receta
            total_venta_libre += sum_venta_libre

        self.write({
            'state': 'step2',
            'preview_line_ids': preview_lines,
            'total_afiliados': len(preview_lines),
            'total_tickets': sum(len(t) for t in tickets_by_uid.values()),
            'total_receta': total_receta,
            'total_venta_libre': total_venta_libre,
            'total_general': total_receta + total_venta_libre,
            'has_overwrites': has_overwrites,
        })

        return self._reopen_wizard()

    def action_confirm_import(self):
        """Step 2 -> Step 3: Create/update records."""
        self.ensure_one()

        try:
            import openpyxl
            from io import BytesIO
        except ImportError:
            raise UserError(_('La librería openpyxl no está instalada.'))

        file_content = base64.b64decode(self.file_data)
        wb = openpyxl.load_workbook(BytesIO(file_content), data_only=True)
        ws = wb.active

        # Re-parse the Excel (we need the full ticket data)
        tickets_by_uid = {}
        for row in ws.iter_rows(min_row=4, max_row=ws.max_row, values_only=False):
            uid_val = row[3].value
            if uid_val is None:
                continue
            try:
                uid = int(uid_val)
            except (ValueError, TypeError):
                continue

            nombre = row[2].value or ''
            fecha_raw = row[5].value
            monto_receta = self._parse_number(row[6].value)
            monto_venta_libre = self._parse_number(row[7].value)

            fecha = None
            if isinstance(fecha_raw, datetime):
                fecha = fecha_raw.date()
            elif isinstance(fecha_raw, str):
                for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
                    try:
                        fecha = datetime.strptime(fecha_raw, fmt).date()
                        break
                    except ValueError:
                        continue

            numero_orden = 0
            try:
                numero_orden = int(row[0].value) if row[0].value else 0
            except (ValueError, TypeError):
                pass

            if uid not in tickets_by_uid:
                tickets_by_uid[uid] = []
            tickets_by_uid[uid].append({
                'uid': uid,
                'nombre': nombre,
                'numero_orden': numero_orden,
                'fecha': fecha,
                'monto_receta': monto_receta,
                'monto_venta_libre': monto_venta_libre,
            })

        Affiliate = self.env['affiliation.affiliate']
        Expense = self.env['affiliate.pharmacy.expenses']
        Ticket = self.env['pharmacy.ticket']

        created_count = 0
        updated_count = 0
        month = self.date_month
        year = int(self.date_year)

        for uid, tickets in tickets_by_uid.items():
            affiliate = Affiliate.search([('uid', '=', uid)], limit=1)
            if not affiliate:
                nombre_ref = tickets[0]['nombre'] if tickets else str(uid)
                raise UserError(_(
                    'El afiliado UID %s (%s) no existe en el sistema.'
                ) % (uid, nombre_ref))

            # Find or create expense record for this affiliate/month/year
            expense = Expense.search([
                ('affiliate_id', '=', affiliate.id),
                ('month', '=', month),
                ('year', '=', year),
            ], limit=1)

            if not expense:
                expense = Expense.create({
                    'affiliate_id': affiliate.id,
                    'month': month,
                    'year': year,
                })

            # Find or create expense line for this pharmacy
            expense_line = expense.linea_gastos_ids.filtered(
                lambda l: l.farmacia_id.id == self.farmacia_id.id
            )

            if not expense_line:
                expense_line = self.env['affiliate.pharmacy.expense.line'].create({
                    'expense_id': expense.id,
                    'farmacia_id': self.farmacia_id.id,
                    'gasto_plan': 0.0,
                    'gasto_venta_libre': 0.0,
                })
                created_count += 1
            else:
                updated_count += 1

            # Delete existing tickets for this line
            expense_line.ticket_ids.unlink()

            # Create new tickets
            sum_receta = 0.0
            sum_venta_libre = 0.0
            for t in tickets:
                Ticket.create({
                    'expense_line_id': expense_line.id,
                    'numero_orden': t['numero_orden'],
                    'nombre_archivo': t['nombre'],
                    'fecha': t['fecha'],
                    'monto_receta': t['monto_receta'],
                    'monto_venta_libre': t['monto_venta_libre'],
                })
                sum_receta += t['monto_receta']
                sum_venta_libre += t['monto_venta_libre']

            # Update expense line totals
            expense_line.write({
                'gasto_plan': sum_receta,
                'gasto_venta_libre': sum_venta_libre,
            })

        total_tickets = sum(len(t) for t in tickets_by_uid.values())
        result_msg = (
            f"Importación completada:\n"
            f"- Registros nuevos: {created_count}\n"
            f"- Registros actualizados: {updated_count}\n"
            f"- Total tickets importados: {total_tickets}"
        )

        self.write({
            'state': 'step3',
            'result_message': result_msg,
        })

        return self._reopen_wizard()

    def action_back(self):
        """Return to step 1."""
        self.ensure_one()
        self.preview_line_ids.unlink()
        self.write({'state': 'step1'})
        return self._reopen_wizard()

    def _reopen_wizard(self):
        """Return action to reopen this wizard."""
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    @staticmethod
    def _parse_number(value):
        """Parse a cell value to float, handling None and strings."""
        if value is None:
            return 0.0
        if isinstance(value, (int, float)):
            return float(value)
        if isinstance(value, str):
            value = value.strip().replace(',', '.')
            try:
                return float(value)
            except ValueError:
                return 0.0
        return 0.0


class PharmacyExpenseImportPreview(models.TransientModel):
    _name = 'pharmacy.expense.import.preview'
    _description = 'Preview de importación de gastos de farmacia'

    wizard_id = fields.Many2one(
        'pharmacy.expense.import.wizard',
        string='Wizard',
        ondelete='cascade',
    )
    affiliate_id = fields.Many2one('affiliation.affiliate', string='Afiliado')
    uid = fields.Integer(string='N° Afiliado')
    nombre_archivo = fields.Char(string='Nombre (archivo)')
    cantidad_tickets = fields.Integer(string='Tickets')
    total_receta = fields.Float(string='Bajo Receta')
    total_venta_libre = fields.Float(string='Venta Libre')
    total = fields.Float(string='Total', compute='_compute_total')
    sobreescribe = fields.Boolean(string='Sobreescribe')

    @api.depends('total_receta', 'total_venta_libre')
    def _compute_total(self):
        for rec in self:
            rec.total = rec.total_receta + rec.total_venta_libre
